# -*- coding: utf-8 -*-
"""Build the Elevora AI final report as an editable .docx.

Typography follows the reference report (finalreport.pdf):
  A4 portrait, 1 inch margins, Times New Roman 12 pt justified body,
  1.5 line spacing, 0.5 inch first-line indent, centred bold 18 pt chapter
  titles, 14 pt section headings, bordered tables and centred figure captions.
"""
import io
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content_part1 as c1
import content_part2 as c2
import content_part3 as c3

TOOLS_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(TOOLS_DIR))     # repository root
REPORT_DIR = os.path.dirname(TOOLS_DIR)                 # report_output/
DOCX = os.path.join(REPORT_DIR, "ElevoraAI_Final_Report.docx")

SERIF = "Times New Roman"
MONO = "Consolas"

# ----------------------------------------------------------------- helpers
TOC_ENTRIES = []   # (kind, key, label)
FIG_ENTRIES = []   # (number, caption)
TBL_ENTRIES = []   # (number, caption)

# --- child element order required by the WordprocessingML schema ----------
PPR_SEQ = ('w:pStyle', 'w:keepNext', 'w:keepLines', 'w:pageBreakBefore', 'w:framePr',
           'w:widowControl', 'w:numPr', 'w:suppressLineNumbers', 'w:pBdr', 'w:shd',
           'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap',
           'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN',
           'w:bidi', 'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind',
           'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
           'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl',
           'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')
RPR_SEQ = ('w:rStyle', 'w:rFonts', 'w:b', 'w:bCs', 'w:i', 'w:iCs', 'w:caps',
           'w:smallCaps', 'w:strike', 'w:dstrike', 'w:outline', 'w:shadow',
           'w:emboss', 'w:imprint', 'w:noProof', 'w:snapToGrid', 'w:vanish',
           'w:webHidden', 'w:color', 'w:spacing', 'w:w', 'w:kern', 'w:position',
           'w:sz', 'w:szCs', 'w:highlight', 'w:u', 'w:effect', 'w:bdr', 'w:shd',
           'w:fitText', 'w:vertAlign', 'w:rtl', 'w:cs', 'w:em', 'w:lang',
           'w:eastAsianLayout', 'w:specVanish', 'w:oMath')
TBLPR_SEQ = ('w:tblStyle', 'w:tblpPr', 'w:tblOverlap', 'w:bidiVisual',
             'w:tblStyleRowBandSize', 'w:tblStyleColBandSize', 'w:tblW', 'w:jc',
             'w:tblCellSpacing', 'w:tblInd', 'w:tblBorders', 'w:shd', 'w:tblLayout',
             'w:tblCellMar', 'w:tblLook', 'w:tblCaption', 'w:tblDescription',
             'w:tblPrChange')
SECTPR_SEQ = ('w:headerReference', 'w:footerReference', 'w:footnotePr', 'w:endnotePr',
              'w:type', 'w:pgSz', 'w:pgMar', 'w:paperSrc', 'w:pgBorders',
              'w:lnNumType', 'w:pgNumType', 'w:cols', 'w:formProt', 'w:vAlign',
              'w:noEndnote', 'w:titlePg', 'w:textDirection', 'w:bidi', 'w:rtlGutter',
              'w:docGrid', 'w:printerSettings', 'w:sectPrChange')


def insert_ordered(parent, element, seq):
    """Insert *element* into *parent* honouring the schema order in *seq*."""
    name = element.tag.split('}')[-1]
    pos = list(seq).index('w:' + name)
    parent.insert_element_before(element, *seq[pos + 1:])
    return element


def split_caption(cap):
    """'Fig. 4.1: Overall architecture' -> ('4.1', 'Overall architecture')."""
    m = re.match(r"^\s*(Fig\.|Figure|Table)\s*([\d.]+)\s*[:.\-]\s*(.*)$", cap)
    if m:
        return m.group(2), m.group(3).strip()
    return "", cap.strip()


def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def set_table_borders(table, sz=6):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:%s" % edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    insert_ordered(tblPr, borders, TBLPR_SEQ)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def cell_margins(table, top=60, bottom=60, left=90, right=90):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:%s" % name)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    insert_ordered(tblPr, mar, TBLPR_SEQ)


def add_field(paragraph, instr, size=11, bold=False):
    run = paragraph.add_run()
    run.font.name = SERIF
    run.font.size = Pt(size)
    run.bold = bold
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run._r.append(fld)
    run2 = paragraph.add_run()
    run2.font.name = SERIF
    run2.font.size = Pt(size)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    run2._r.append(it)
    run3 = paragraph.add_run()
    run3.font.name = SERIF
    run3.font.size = Pt(size)
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run3._r.append(fld2)


def add_border(paragraph, edge="bottom", sz=12, space=4):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        insert_ordered(pPr, bdr, PPR_SEQ)
    el = OxmlElement("w:%s" % edge)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), "000000")
    bdr.append(el)


def set_pgnum(section, fmt, start=None):
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(old)
    el = OxmlElement("w:pgNumType")
    el.set(qn("w:fmt"), fmt)
    if start is not None:
        el.set(qn("w:start"), str(start))
    insert_ordered(sectPr, el, SECTPR_SEQ)


def style_run(run, size=12, bold=False, italic=False, font=SERIF, caps=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), font)
    if caps:
        c = OxmlElement("w:caps")
        c.set(qn("w:val"), "1")
        insert_ordered(rPr, c, RPR_SEQ)
    return run


# ----------------------------------------------------------------- document
def build():
    doc = Document()

    # base style
    normal = doc.styles["Normal"]
    normal.font.name = SERIF
    normal.font.size = Pt(12)
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), SERIF)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    pf.first_line_indent = Inches(0.5)

    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(0.9)
    sec.header_distance = Inches(0.45)
    sec.footer_distance = Inches(0.45)

    # ---------------- front matter footer (roman numerals, centred)
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.first_line_indent = Inches(0)
    fp.paragraph_format.space_before = Pt(0)
    add_field(fp, " PAGE ", size=10)
    set_pgnum(sec, "lowerRoman", 1)

    # ---------------- styles used by headings
    styles = doc.styles
    ch_style = styles.add_style("ChapterTitle", 1)   # 1 = paragraph style
    ch_style.base_style = styles["Normal"]
    ch_style.font.name = SERIF
    ch_style.font.size = Pt(18)
    ch_style.font.bold = True
    ch_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ch_style.paragraph_format.space_before = Pt(6)
    ch_style.paragraph_format.space_after = Pt(18)
    ch_style.paragraph_format.first_line_indent = Inches(0)
    ch_style.paragraph_format.keep_with_next = True

    sec_style = styles.add_style("SectionHeading", 1)
    sec_style.base_style = styles["Normal"]
    sec_style.font.name = SERIF
    sec_style.font.size = Pt(14)
    sec_style.font.bold = True
    sec_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sec_style.paragraph_format.space_before = Pt(14)
    sec_style.paragraph_format.space_after = Pt(6)
    sec_style.paragraph_format.first_line_indent = Inches(0)
    sec_style.paragraph_format.keep_with_next = True

    cap_style = styles.add_style("FigureCaption", 1)
    cap_style.base_style = styles["Normal"]
    cap_style.font.name = SERIF
    cap_style.font.size = Pt(12)
    cap_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_style.paragraph_format.first_line_indent = Inches(0)
    cap_style.paragraph_format.space_before = Pt(4)
    cap_style.paragraph_format.space_after = Pt(14)
    cap_style.paragraph_format.keep_with_next = False

    def para(text, size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             first_indent=0.5, space_before=0, space_after=0, line=1.5, font=SERIF,
             keep_with_next=False):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = align
        pf.first_line_indent = Inches(first_indent)
        pf.space_before = Pt(space_before)
        pf.space_after = Pt(space_after)
        pf.line_spacing = line
        pf.keep_with_next = keep_with_next
        if text:
            style_run(p.add_run(text), size=size, bold=bold, italic=italic, font=font)
        return p

    def bullet(text, numbered=False, number=None):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.left_indent = Inches(0.5)
        pf.first_line_indent = Inches(-0.25)
        pf.space_after = Pt(4)
        pf.line_spacing = 1.4
        marker = ("%s. " % number) if numbered else ("•  ")
        style_run(p.add_run(marker), size=12)
        style_run(p.add_run(text), size=12)
        return p

    def chapter(number_text, title_text, key):
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.keep_with_next = True
        style_run(p.add_run(number_text), size=18, bold=True)
        tp = doc.add_paragraph(style="ChapterTitle")
        style_run(tp.add_run(title_text), size=18, bold=True)
        TOC_ENTRIES.append(("chapter", key, title_text))

    def section(number_text, title_text):
        p = doc.add_paragraph(style="SectionHeading")
        if title_text is None:
            # an unnumbered, labelled subsection such as "Future Enhancements"
            text = number_text
            key, label = number_text, number_text
        else:
            text = "%s  %s" % (number_text, title_text)
            key, label = number_text, title_text
        style_run(p.add_run(text), size=14, bold=True)
        TOC_ENTRIES.append(("section", key, label))
        return p

    def table_block(caption, headers, rows, widths, number):
        cp = doc.add_paragraph()
        cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.first_line_indent = Inches(0)
        cp.paragraph_format.space_before = Pt(10)
        cp.paragraph_format.space_after = Pt(4)
        cp.paragraph_format.keep_with_next = True
        cp.paragraph_format.line_spacing = 1.0
        style_run(cp.add_run("Table %s: %s" % (number, split_caption(caption)[1])), size=11, bold=True)
        TBL_ENTRIES.append((number, "Table %s: %s" % (number, split_caption(caption)[1])))

        # scale the requested column widths so that the table never exceeds
        # the printable width of the page (A4 minus the two 1 inch margins)
        scale = 6.2 / float(sum(widths))
        widths = [w * scale for w in widths]

        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        set_table_borders(t)
        cell_margins(t)
        hdr = t.rows[0]
        set_repeat_header(hdr)
        for i, h in enumerate(headers):
            c = hdr.cells[i]
            c.text = ""
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            style_run(p.add_run(h), size=10.5, bold=True)
            set_cell_bg(c, "E8E8E8")
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                c = cells[i]
                c.text = ""
                p = c.paragraphs[0]
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i else WD_ALIGN_PARAGRAPH.CENTER
                style_run(p.add_run(str(val)), size=10)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    def figure_block(rel_path, caption, width, number):
        path = os.path.join(REPORT_DIR, rel_path)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Inches(width))
        cp = doc.add_paragraph(style="FigureCaption")
        cap_text = "Fig. %s: %s" % (number, split_caption(caption)[1])
        style_run(cp.add_run(cap_text), size=12)
        FIG_ENTRIES.append((number, cap_text))

    def code_block(source, text):
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t, sz=4)
        cell_margins(t, top=80, bottom=80, left=120, right=120)
        c = t.rows[0].cells[0]
        c.text = ""
        set_cell_bg(c, "F4F4F4")
        first = True
        for line in text.split("\n"):
            p = c.paragraphs[0] if first else c.add_paragraph()
            first = False
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style_run(p.add_run(line), size=9, font=MONO)
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.first_line_indent = Inches(0)
        cap.paragraph_format.space_before = Pt(3)
        cap.paragraph_format.space_after = Pt(12)
        cap.paragraph_format.line_spacing = 1.0
        style_run(cap.add_run("Extract: %s" % source), size=10, italic=True)

    # --------------------------------------------------------- render blocks
    def render(blocks):
        for blk in blocks:
            kind = blk[0]
            if kind == "h1":
                if blk[1] == "ABSTRACT":
                    in_abstract[0] = True
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.first_line_indent = Inches(0)
                    p.paragraph_format.space_after = Pt(16)
                    p.paragraph_format.line_spacing = 1.15
                    style_run(p.add_run("ABSTRACT"), size=18, bold=True)
                elif blk[1] == "CONCLUSION":
                    doc.add_page_break()
                    p = doc.add_paragraph(style="ChapterTitle")
                    p.paragraph_format.space_after = Pt(16)
                    style_run(p.add_run("CONCLUSION"), size=18, bold=True)
                    TOC_ENTRIES.append(("chapter", "CONCLUSION", "CONCLUSION"))
                elif blk[1] == "REFERENCES":
                    doc.add_page_break()
                    p = doc.add_paragraph(style="ChapterTitle")
                    p.paragraph_format.space_after = Pt(16)
                    style_run(p.add_run("REFERENCES"), size=18, bold=True)
                    TOC_ENTRIES.append(("chapter", "REFERENCES", "REFERENCES"))
                else:
                    # the body already starts on a new page because of the
                    # section break, so no extra page break is needed there
                    if blk[1] != "CHAPTER – 1":
                        doc.add_page_break()
                    chapter(blk[1], blk[2], blk[1])
            elif kind == "h2":
                section(blk[1], blk[2])
            elif kind == "p":
                para(blk[1])
            elif kind == "b":
                bullet(blk[1])
            elif kind == "n":
                bullet(blk[1], numbered=True, number=blk[2] if len(blk) > 2 else "")
            elif kind == "table":
                m = re.search(r"Table\s+(\d+\.\d+)", blk[1])
                if m:
                    number = m.group(1)
                else:
                    table_no[0] += 1
                    number = "%d.%d" % (chapter_no[0], table_no[0])
                table_block(blk[1], blk[2], blk[3], blk[4], number)
            elif kind == "fig":
                m = re.search(r"(?:Fig\.|Figure)\s*(\d+\.\d+)", blk[2])
                if m:
                    number = m.group(1)
                else:
                    fig_no[0] += 1
                    number = "%d.%d" % (chapter_no[0], fig_no[0])
                figure_block(blk[1], blk[2], blk[3], number)
            elif kind == "code":
                code_block(blk[1], blk[2])
            elif kind == "ref":
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                pf.left_indent = Inches(0.4)
                pf.first_line_indent = Inches(-0.4)
                pf.space_after = Pt(6)
                pf.line_spacing = 1.25
                style_run(p.add_run(blk[1] + " "), size=11, bold=True)
                style_run(p.add_run(blk[2] + " "), size=11, italic=True)
                style_run(p.add_run("\"" + blk[3] + "\", "), size=11)
                style_run(p.add_run(blk[4]), size=11)

    # ----------------------------------------------------------- front matter
    chapter_no = [0]
    table_no = [0]
    fig_no = [0]

    render(c1.FRONT)

    # ---- TABLE OF CONTENTS placeholder table
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(14)
    style_run(p.add_run("TABLE OF CONTENTS"), size=18, bold=True)

    # ---------------- end of front matter: insert section break for body
    body_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    body_sec.page_width = Inches(8.27)
    body_sec.page_height = Inches(11.69)
    body_sec.left_margin = Inches(1.0)
    body_sec.right_margin = Inches(1.0)
    body_sec.top_margin = Inches(1.0)
    body_sec.bottom_margin = Inches(0.9)
    body_sec.header_distance = Inches(0.45)
    body_sec.footer_distance = Inches(0.45)
    set_pgnum(body_sec, "decimal", 1)

    body_sec.header.is_linked_to_previous = False
    body_sec.footer.is_linked_to_previous = False

    hp = body_sec.header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.first_line_indent = Inches(0)
    hp.paragraph_format.space_after = Pt(2)
    hp.paragraph_format.line_spacing = 1.0
    tabs = hp.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(6.27), WD_TAB_ALIGNMENT.RIGHT)
    style_run(hp.add_run("Elevora AI  -  Dev-Elevate"), size=11)
    style_run(hp.add_run("\t"), size=11)
    add_field(hp, ' STYLEREF "ChapterTitle" \\* MERGEFORMAT ', size=11)
    add_border(hp, "bottom", sz=18, space=3)

    ftr = body_sec.footer.paragraphs[0]
    ftr.text = ""
    ftr.paragraph_format.first_line_indent = Inches(0)
    ftr.paragraph_format.space_before = Pt(2)
    ftr.paragraph_format.line_spacing = 1.0
    tabs = ftr.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(3.13), WD_TAB_ALIGNMENT.CENTER)
    tabs.add_tab_stop(Inches(6.27), WD_TAB_ALIGNMENT.RIGHT)
    style_run(ftr.add_run("Elevora AI"), size=11)
    style_run(ftr.add_run("\t"), size=11)
    style_run(ftr.add_run("Dev-Elevate"), size=11)
    style_run(ftr.add_run("\t"), size=11)
    add_field(ftr, " PAGE ", size=11)
    add_border(ftr, "top", sz=6, space=3)

    # ----------------------------------------------------------- body
    chapter_no[0] = 1
    render(c1.CH1)
    chapter_no[0] = 2
    render(c1.CH2)
    chapter_no[0] = 3
    render(c2.CH3)
    chapter_no[0] = 4
    render(c2.CH4)
    chapter_no[0] = 5
    render(c3.CH5)
    chapter_no[0] = 6
    render(c3.CH6)
    render(c3.CONCLUSION)
    render(c3.REFERENCES)

    # ------------------------------------------- build contents/figure tables
    build_front_matter_tables(doc)

    doc.save(DOCX)
    dump_entries()
    print("saved", DOCX)
    print("toc entries", len(TOC_ENTRIES), "figures", len(FIG_ENTRIES),
          "tables", len(TBL_ENTRIES))


def dump_entries():
    """Write the TOC / figure / table keys so the page resolver can use them."""
    import json
    data = {
        "toc": [{"kind": k, "key": key, "label": label} for k, key, label in TOC_ENTRIES],
        "figures": [{"number": n, "caption": c} for n, c in FIG_ENTRIES],
        "tables": [{"number": n, "caption": c} for n, c in TBL_ENTRIES],
    }
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "toc_entries.json")
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=1, ensure_ascii=False)
    print("wrote", path)


# ------------------------------------------------------------------ front matter tables
PAGEMAP_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "pagemap.json")


def load_pagemap():
    """Page numbers resolved from the previous PDF render (key -> page label)."""
    try:
        import json
        with open(PAGEMAP_PATH, encoding="utf-8") as fh:
            data = json.load(fh)
        if isinstance(data, dict):
            return data
    except Exception:
        pass
    return {}


def clean_caption(cap):
    """'Fig. 4.1: Overall architecture' -> 'Overall architecture'."""
    import re
    s = re.sub(r"^\s*(Fig\.|Figure|Table)\s*[\d.]+\s*[:.\-]?\s*", "", cap)
    return s.strip().rstrip(".")


def build_front_matter_tables(doc):
    """Insert the TOC, list of figures and list of tables in the front matter."""
    pagemap = load_pagemap()

    # locate the paragraph that holds the "TABLE OF CONTENTS" heading
    toc_p = None
    for p in doc.paragraphs:
        if p.text.strip() == "TABLE OF CONTENTS":
            toc_p = p
            break
    assert toc_p is not None

    def new_table(headers, widths, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        set_table_borders(t)
        cell_margins(t, top=50, bottom=50, left=90, right=90)
        hdr = t.rows[0]
        set_repeat_header(hdr)
        for i, h in enumerate(headers):
            c = hdr.cells[i]
            c.text = ""
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            style_run(p.add_run(h), size=12, bold=True)
            set_cell_bg(c, "E8E8E8")
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                c = cells[i]
                c.text = ""
                p = c.paragraphs[0]
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                if i == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    style_run(p.add_run(str(val)), size=11, bold=True)
                elif i == len(headers) - 1:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    style_run(p.add_run(str(val)), size=11)
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    style_run(p.add_run(str(val)), size=11)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
        return t

    def new_heading(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_after = Pt(14)
        p.paragraph_format.space_before = Pt(0)
        style_run(p.add_run(text), size=18, bold=True)
        return p

    # ---- contents rows -------------------------------------------------
    rows = []
    sl = 0
    for kind, key, label in TOC_ENTRIES:
        page = pagemap.get(key, "00")
        if kind == "chapter":
            if key.startswith("CHAPTER"):
                sl += 1
                rows.append((sl, label, page))
            else:
                rows.append(("", label, page))
        else:
            if key == label:
                rows.append(("", label, page))
            else:
                rows.append(("", "%s   %s" % (key, label), page))

    toc_tbl = new_table(["SL. NO", "NAME", "PAGE NO"], [0.9, 4.1, 1.27], rows)

    # ---- list of figures ------------------------------------------------
    fig_rows = [(num, clean_caption(cap), pagemap.get("FIG:" + num, "00"))
                for num, cap in FIG_ENTRIES]

    # ---- list of tables -------------------------------------------------
    tbl_rows = [(num, clean_caption(cap), pagemap.get("TBL:" + num, "00"))
                for num, cap in TBL_ENTRIES]

    fig_p = new_heading("FIGURE OF CONTENT")
    fig_tbl = new_table(["FIG NO", "NAME", "PAGE NO"], [0.9, 4.1, 1.27], fig_rows)
    tbl_p = new_heading("LIST OF TABLES")
    tbl_tbl = new_table(["TABLE NO", "NAME", "PAGE NO"], [0.9, 4.1, 1.27], tbl_rows)

    # ---- move everything into the front matter, in the intended order ---
    # anchor starts at the TABLE OF CONTENTS heading; every element is
    # inserted directly after the previous one so the order is preserved.
    anchor = toc_p._p
    for el in (toc_tbl._tbl, fig_p._p, fig_tbl._tbl, tbl_p._p, tbl_tbl._tbl):
        anchor.addnext(el)
        anchor = el

    # the two list headings start on a fresh page
    for p in (fig_p, tbl_p):
        pb = OxmlElement("w:pageBreakBefore")
        insert_ordered(p._p.get_or_add_pPr(), pb, PPR_SEQ)

    return toc_tbl, fig_tbl, tbl_tbl


if __name__ == "__main__":
    build()
