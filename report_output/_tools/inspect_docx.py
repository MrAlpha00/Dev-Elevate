# -*- coding: utf-8 -*-
"""Structural sanity check of the generated DOCX."""
import io
import os
import sys

from docx import Document
from docx.oxml.ns import qn

REPORT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX = os.path.join(REPORT_DIR, "ElevoraAI_Final_Report.docx")

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

doc = Document(DOCX)
body = doc.element.body

print("sections:", len(doc.sections))
for i, s in enumerate(doc.sections):
    pgn = s._sectPr.find(qn("w:pgNumType"))
    fmt = pgn.get(qn("w:fmt")) if pgn is not None else None
    start = pgn.get(qn("w:start")) if pgn is not None else None
    print("  section %d: %.2fx%.2f in, margins L%.2f R%.2f T%.2f B%.2f, pgnum fmt=%s start=%s" % (
        i, s.page_width.inches, s.page_height.inches, s.left_margin.inches,
        s.right_margin.inches, s.top_margin.inches, s.bottom_margin.inches, fmt, start))
    print("    header linked:", s.header.is_linked_to_previous, "text:",
          " | ".join(p.text for p in s.header.paragraphs))
    print("    footer linked:", s.footer.is_linked_to_previous, "text:",
          " | ".join(p.text for p in s.footer.paragraphs))

print("\nbody children (first 30):")
for i, el in enumerate(body):
    tag = el.tag.split("}")[-1]
    txt = ""
    if tag == "p":
        txt = "".join(el.itertext())[:70]
    elif tag == "tbl":
        txt = "TABLE"
    elif tag == "sectPr":
        txt = "SECTPR(final)"
    if i < 30:
        print("  %3d %-7s %s" % (i, tag, txt))

# locate the section break paragraph
print("\nsection-break paragraph position:")
for i, el in enumerate(body):
    if el.tag.endswith("}p"):
        pPr = el.find(qn("w:pPr"))
        if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
            print("  index", i, repr("".join(el.itertext())[:60]))

print("\ntables:", len(doc.tables))
for i, t in enumerate(doc.tables[:6]):
    print("  table %d rows=%d cols=%d first=%r" % (
        i, len(t.rows), len(t.columns), t.rows[0].cells[0].text[:30]))
